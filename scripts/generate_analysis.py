import sqlite3
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

db_path = r'F:\Mydownloads\xtcz10\switch.db'
conn = sqlite3.connect(db_path)
conn.row_factory = sqlite3.Row
cur = conn.cursor()

cur.execute("SELECT id, module, display, serverId, tips, extra FROM module_switch ORDER BY module")
rows = cur.fetchall()
conn.close()

# Function inference based on extra content
def infer_function(module, display, extra):
    if not extra or extra.strip() in ('', '无'):
        return '配置型开关（无额外参数，需结合源码确认）', '低'
    
    e = extra.strip()
    confidence = '高'
    
    # Duration
    if '"duration":30000' in e:
        return '定时/持续操作开关（30秒时长控制，如动画、提示、超时）', '高'
    
    # Night WiFi strategy
    if 'nightStartHour' in e and 'wifiNum' in e:
        return '夜间WiFi省电策略（23:00-6:00，按步长50扫描，每次10个热点）', '高'
    
    # Data trace
    if 'deadLineTime' in e and 'dropCount' in e and 'trace' in e:
        return '数据追踪/上报开关（3天有效期，丢弃阈值90，52ms采样间隔）', '高'
    
    # WiFi scan
    if 'scanTime' in e and 'scanChanel' in e:
        return 'WiFi信道扫描配置（40秒扫描周期，覆盖1-12信道）', '高'
    
    # Flip wrist
    if 'flipCount' in e and 'flipDuration' in e:
        return '翻腕亮屏控制（60秒内最多10次，超限禁用180秒）', '高'
    
    # Max valid time
    if 'maxValidTime' in e:
        return '数据有效期控制（最长7天有效）', '高'
    
    # Friend filter
    if 'filterModel' in e and 'friendType' in e:
        models = e.count('","') + 1 if 'filterModel' in e else 0
        return f'好友/联系人型号过滤（支持指定型号，好友类型位掩码）', '高'
    
    # Quick setting
    if 'switchId' in e and 'isQuickSetting' in e:
        qs = '快速设置' if '1' in e.split('isQuickSetting')[1][:3] else '普通设置'
        ds = e.split('defaultStatus')[1][:5] if 'defaultStatus' in e else ''
        return f'{qs}开关项（默认状态{"开" if "1" in ds else "关"}，可选择）', '高'
    
    # Status
    if e == '{"status" : 1}' or e == '{"status":1}':
        return '状态开关（启用状态=1）', '中'
    
    # Support model
    if 'supportModel' in e:
        return '设备型号支持列表（限定可使用该功能的手表型号）', '高'
    
    # Max count
    if 'maxCount' in e and '299' in e:
        return '数量上限控制（最大299条/个）', '高'
    
    # Domain whitelist
    if 'domainList' in e:
        return '网络域名白名单（40+教育/支付/电商域名，用于网络访问控制）', '高'
    
    # Bubble/balance game
    if 'bubbles1' in e or 'balanceball' in e or 'momentmood' in e:
        return '互动游戏防沉迷限制（气泡点击60次/10分，平衡球10分/次，弹幕10条）', '高'
    
    # Limit count
    if 'limitCount' in e:
        return '数量限制开关（限制数量=5）', '高'
    
    # MAC prefix
    if 'prefix' in e and ':' in e:
        return 'BLE/蓝牙设备MAC前缀匹配（用于识别特定外设）', '高'
    
    # Message send control
    if 'allowCount' in e and 'delaySendTime' in e:
        return '消息发送风控（允许1000条，延迟30秒发送，违规记录24小时）', '高'
    
    # Map provider
    if '"baidu"' in e and '"gaode"' in e and '"tengxun"' in e:
        baidu = 'true' in e.split('"baidu"')[1][:8]
        gaode = 'true' in e.split('"gaode"')[1][:8]
        tengxun = 'true' in e.split('"tengxun"')[1][:8]
        providers = []
        if baidu: providers.append('百度')
        if gaode: providers.append('高德')
        if tengxun: providers.append('腾讯')
        return f'地图供应商选择（{"+".join(providers) if providers else "无"}）', '高'
    
    # Device regex
    if 'regexs' in e and ('RMTT' in e or 'TELLO' in e or 'JimuGo' in e):
        return '外设设备正则匹配（大疆RMTT/TELLO无人机、JimuGo积木机器人）', '高'
    
    # Event tracking
    if 'videocall_result' in e or 'BigDataSDK' in e or 'SST_QCRIL' in e:
        return '埋点/日志事件采集配置（20个事件：视频通话、重启监控、高通基带异常、应用启动、定位、OCR翻译等）', '高'
    
    # Reboot type
    if 'reboottype' in e:
        return '重启类型配置（reboottype=0）', '中'
    
    # User debug
    if 'userdebug' in e:
        return '用户调试模式标记（userdebug=false，正式版）', '高'
    
    # Auto scroll
    if 'AutoScrollTime' in e:
        return '自动滚动控制（5秒自动滚动，允许20秒滚动周期）', '高'
    
    # Exit time
    if 'exitTime' in e:
        return '退出超时控制（180秒无操作自动退出）', '高'
    
    # BGM
    if 'shineeBgmId' in e or 'creationBgmId' in e:
        return '背景音乐ID配置（闪亮模式5首，创作模式10首）', '高'
    
    # Real-time location GIF
    if 'RealTimeLocationGifVaildTime' in e:
        return '实时定位GIF有效期控制（时间戳校验）', '高'
    
    # Mode
    if e == '{ "mode":1}' or e == '{"mode":1}':
        return '模式选择开关（mode=1）', '中'
    
    # NSFW content filter
    if 'nsfwPornWeights' in e or 'nsfwWeights' in e or 'supportingVideo' in e:
        return '内容安全/NSFW审核（支持视频+GIF，视频帧5秒/GIF帧1秒，涉黄权重阈值0.2）', '高'
    
    # Vacation
    if 'winterVacationStart' in e or 'summerVacationStart' in e:
        return '假期模式/异常检测（寒暑假日期配置，相似度/电池/退出阈值）', '高'
    
    # Accuracy
    if 'accuracy' in e and '15' in e:
        return '定位精度控制（精度阈值=15米）', '高'
    
    # Album num
    if 'album_num' in e:
        return '相册数量控制（album_num=5）', '高'
    
    # Sync
    if 'SendSyncRequestInterval' in e:
        return '数据同步配置（120秒同步间隔，20秒超时）', '高'
    
    # Hide items
    if 'hideItems' in e:
        items = e.split('[')[1].split(']')[0] if '[' in e else ''
        return f'界面项目隐藏（隐藏项：[{items}]）', '高'
    
    # Scheduled reboot
    if 'timeIntervalList' in e and 'bootTimeInHours' in e:
        return '定时重启策略（1:00-5:00窗口，开机48小时后触发，延迟240分钟，重启Launcher）', '高'
    
    # Sharing finish
    if 'sharingFinishTime' in e:
        return '分享完成超时（120秒）', '高'
    
    # Pure number
    if e.strip().isdigit():
        return f'数值配置开关（值={e.strip()}）', '中'
    
    # Sport types
    if 'sportTypeBicycling' in e or 'sportTypeWalk' in e:
        return '运动模式支持列表（11种：骑行/登山/步行/游泳/滑雪/篮球/燃脂/乒乓球/滑板/羽毛球/街舞）', '高'
    
    # Group chat limit
    if 'groupChatDailyLimit' in e:
        return '群聊每日限制（每日1次）', '高'
    
    # Type max count
    if '"type":0' in e and '"maxCount"' in e:
        return '消息/表情数量限制（type0和type1各最多20条）', '高'
    
    # Correct distance
    if 'vaildCorrectDistance' in e:
        return '定位有效校正距离（70米内校正有效）', '高'
    
    # Array config
    if e.strip().startswith('[') and e.strip().endswith(']') and len(e) < 20:
        return f'数组配置开关（值={e.strip()}）', '中'
    
    # Crash recovery
    if 'crashTimes' in e and 'reBootEnable' in e:
        return '崩溃自愈/恢复策略（3次崩溃7.5分钟内触发，可重启/清数据，24小时窗口）', '高'
    
    # Task cost time
    if 'task_cost_time' in e:
        return '任务耗时阈值（2000ms）', '高'
    
    # Web URL
    if 'web_url' in e and 'http' in e:
        return 'Web页面加载配置（指定H5页面URL：ad2022.html）', '高'
    
    # Log config
    if 'maxSize' in e and 'maxBackupIndex' in e:
        return '日志轮转配置（单文件512KB，保留4个备份，最长7天）', '高'
    
    # Heatmap
    if 'mColors' in e and 'mRadiusMeter' in e:
        return '定位热力图配置（4色渐变，半径15米，最大强度150，12小时有效期）', '高'
    
    # Egg animation
    if 'playEggAnimationLimit' in e:
        return '彩蛋动画限制（单次6个，每日40个）', '高'
    
    # Low power/wechat
    if 'lowPowerThreshold' in e and 'wechatSwitch' in e:
        return '低电量/微信联动（低于10%触发，微信开关000，不自动显示）', '高'
    
    # Max temp
    if 'maxAccurateTemp' in e:
        return '体温测量上限（最高精确测温42°C）', '高'
    
    # Shutdown battery
    if 'shutdown_battery' in e and 'restart_time' in e:
        return '低电量自动关机（8%电量触发，9秒倒计时，60秒后可重启）', '高'
    
    # Effective time
    if 'effectiveTime' in e and '43200' in e:
        return '功能有效期控制（43200秒=12小时）', '高'
    
    # Storage limit
    if 'storageLimit' in e:
        return '存储容量限制（按型号配置：多数100MB，D3为50MB，默认50MB）', '高'
    
    # Max duration
    if 'max_duration' in e:
        return '最大时长限制（600秒=10分钟）', '高'
    
    # Alipay red dot
    if 'isShowRedDot' in e and 'alipay' in e:
        return '支付宝红点提示（显示红点，跳转支付宝加载页）', '高'
    
    # Interval trigger
    if '"interval"' in e and 'triggerCount' in e:
        return '周期触发配置（24小时间隔，触发3次）', '高'
    
    # Network reconnect
    if 'triggerNetPeriod' in e or 'imConnectPeriod' in e:
        return '网络/IM重连策略（5次5秒间隔后改为10秒，不主动触发网络）', '高'
    
    # Open retry
    if 'open_retry' in e:
        return '打开重试机制（重试1次）', '高'
    
    return f'扩展配置开关（参数：{e[:80]}）', '中'


# Create workbook
wb = Workbook()

# Sheet 1: Overview
ws1 = wb.active
ws1.title = '数据库概览'

header_font = Font(bold=True, size=12, color='FFFFFF')
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

overview_data = [
    ['项目', '值', '说明'],
    ['文件名', 'switch.db', 'SQLite 3 格式数据库'],
    ['文件大小', '53,248 字节 (52KB)', ''],
    ['语言环境', 'zh_CN', '简体中文'],
    ['表数量', '3', 'android_metadata, module_switch, sqlite_sequence'],
    ['核心表', 'module_switch', '模块开关配置表'],
    ['模块总数', '434', '条模块开关记录'],
    ['display=1（启用/显示）', '56', '占比 12.9%'],
    ['display=0（隐藏/关闭）', '378', '占比 87.1%'],
    ['有extra配置', '81', '占比 18.7%，含JSON参数'],
    ['无extra配置', '353', '占比 81.3%，纯开关型'],
    ['tips非空', '0', '提示文案字段全部为空'],
    ['id范围', '55069 - 55502', '自增主键'],
    ['module范围', '3 - 3415', '模块代码（唯一约束）'],
    ['自增序列当前值', '55502', 'sqlite_sequence.seq'],
    ['关联产品', '小天才/xtc 儿童智能手表', 'ND03 V2.8.1 构建'],
    ['支持设备型号', '17+种', 'I16/I17/I20/I25/I32/ND01/ND03/D3等'],
]

for r, row in enumerate(overview_data, 1):
    for c, val in enumerate(row, 1):
        cell = ws1.cell(row=r, column=c, value=val)
        cell.border = thin_border
        if r == 1:
            cell.font = header_font
            cell.fill = header_fill
        cell.alignment = Alignment(vertical='center', wrap_text=True)

ws1.column_dimensions['A'].width = 28
ws1.column_dimensions['B'].width = 35
ws1.column_dimensions['C'].width = 50

# Sheet 2: Field description
ws2 = wb.create_sheet('字段说明')
field_data = [
    ['字段名', '类型', '说明'],
    ['id', 'INTEGER (PK, AUTOINCREMENT)', '自增主键，记录唯一标识，范围55069-55502'],
    ['module', 'INTEGER (UNIQUE)', '模块代码，功能模块的唯一编号，范围3-3415，不连续'],
    ['display', 'INTEGER', '显示/启用状态：0=隐藏/关闭，1=显示/开启'],
    ['serverId', 'INTEGER', '服务端配置ID，对应后台管理系统中的配置记录ID'],
    ['tips', 'VARCHAR', '提示文案，当前全部为空（NULL）'],
    ['extra', 'VARCHAR', '扩展配置，JSON格式字符串，存储该模块的具体参数（81条有值）'],
]
for r, row in enumerate(field_data, 1):
    for c, val in enumerate(row, 1):
        cell = ws2.cell(row=r, column=c, value=val)
        cell.border = thin_border
        if r == 1:
            cell.font = header_font
            cell.fill = header_fill
        cell.alignment = Alignment(vertical='center', wrap_text=True)
ws2.column_dimensions['A'].width = 15
ws2.column_dimensions['B'].width = 30
ws2.column_dimensions['C'].width = 70

# Sheet 3: Full module analysis
ws3 = wb.create_sheet('全部模块功能分析')
full_header = ['序号', 'module代码', 'display', 'serverId', '有extra', '功能推断', '推断置信度', 'extra配置详情']
for c, val in enumerate(full_header, 1):
    cell = ws3.cell(row=1, column=c, value=val)
    cell.font = header_font
    cell.fill = header_fill
    cell.border = thin_border
    cell.alignment = Alignment(horizontal='center', vertical='center')

enabled_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
extra_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')

for idx, r in enumerate(rows, 1):
    module = r['module']
    display = r['display']
    serverId = r['serverId']
    extra = r['extra']
    has_extra = '是' if (extra and extra.strip() not in ('', '无')) else '否'
    func, conf = infer_function(module, display, extra)
    
    row_num = idx + 1
    ws3.cell(row=row_num, column=1, value=idx)
    ws3.cell(row=row_num, column=2, value=module)
    ws3.cell(row=row_num, column=3, value=display)
    ws3.cell(row=row_num, column=4, value=serverId)
    ws3.cell(row=row_num, column=5, value=has_extra)
    ws3.cell(row=row_num, column=6, value=func)
    ws3.cell(row=row_num, column=7, value=conf)
    ws3.cell(row=row_num, column=8, value=extra if extra else '')
    
    for c in range(1, 9):
        cell = ws3.cell(row=row_num, column=c)
        cell.border = thin_border
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        if display == 1:
            cell.fill = enabled_fill
        elif has_extra == '是':
            cell.fill = extra_fill

ws3.column_dimensions['A'].width = 6
ws3.column_dimensions['B'].width = 12
ws3.column_dimensions['C'].width = 8
ws3.column_dimensions['D'].width = 10
ws3.column_dimensions['E'].width = 8
ws3.column_dimensions['F'].width = 55
ws3.column_dimensions['G'].width = 10
ws3.column_dimensions['H'].width = 80
ws3.freeze_panes = 'A2'

# Sheet 4: Modules with extra (detailed)
ws4 = wb.create_sheet('带extra配置模块详解')
detail_header = ['module代码', 'display', 'serverId', '功能推断', '完整extra配置']
for c, val in enumerate(detail_header, 1):
    cell = ws4.cell(row=1, column=c, value=val)
    cell.font = header_font
    cell.fill = header_fill
    cell.border = thin_border
    cell.alignment = Alignment(horizontal='center', vertical='center')

extra_rows = [r for r in rows if r['extra'] and r['extra'].strip() not in ('', '无')]
for idx, r in enumerate(extra_rows, 1):
    func, conf = infer_function(r['module'], r['display'], r['extra'])
    row_num = idx + 1
    ws4.cell(row=row_num, column=1, value=r['module'])
    ws4.cell(row=row_num, column=2, value=r['display'])
    ws4.cell(row=row_num, column=3, value=r['serverId'])
    ws4.cell(row=row_num, column=4, value=func)
    ws4.cell(row=row_num, column=5, value=r['extra'])
    for c in range(1, 6):
        cell = ws4.cell(row=row_num, column=c)
        cell.border = thin_border
        cell.alignment = Alignment(vertical='center', wrap_text=True)

ws4.column_dimensions['A'].width = 12
ws4.column_dimensions['B'].width = 8
ws4.column_dimensions['C'].width = 10
ws4.column_dimensions['D'].width = 55
ws4.column_dimensions['E'].width = 100
ws4.freeze_panes = 'A2'

# Sheet 5: Enabled modules
ws5 = wb.create_sheet('已启用模块(display=1)')
enabled_header = ['序号', 'module代码', 'serverId', '有extra', '功能推断', 'extra配置']
for c, val in enumerate(enabled_header, 1):
    cell = ws5.cell(row=1, column=c, value=val)
    cell.font = header_font
    cell.fill = header_fill
    cell.border = thin_border
    cell.alignment = Alignment(horizontal='center', vertical='center')

enabled_rows = [r for r in rows if r['display'] == 1]
for idx, r in enumerate(enabled_rows, 1):
    func, conf = infer_function(r['module'], r['display'], r['extra'])
    has_extra = '是' if (r['extra'] and r['extra'].strip() not in ('', '无')) else '否'
    row_num = idx + 1
    ws5.cell(row=row_num, column=1, value=idx)
    ws5.cell(row=row_num, column=2, value=r['module'])
    ws5.cell(row=row_num, column=3, value=r['serverId'])
    ws5.cell(row=row_num, column=4, value=has_extra)
    ws5.cell(row=row_num, column=5, value=func)
    ws5.cell(row=row_num, column=6, value=r['extra'] if r['extra'] else '')
    for c in range(1, 7):
        cell = ws5.cell(row=row_num, column=c)
        cell.border = thin_border
        cell.alignment = Alignment(vertical='center', wrap_text=True)

ws5.column_dimensions['A'].width = 6
ws5.column_dimensions['B'].width = 12
ws5.column_dimensions['C'].width = 10
ws5.column_dimensions['D'].width = 8
ws5.column_dimensions['E'].width = 55
ws5.column_dimensions['F'].width = 80
ws5.freeze_panes = 'A2'

output_path = r'H:\xtcz10check\switch_db_模块功能分析.xlsx'
wb.save(output_path)
print(f'Excel saved: {output_path}')
print(f'Total modules: {len(rows)}')
print(f'Modules with extra: {len(extra_rows)}')
print(f'Enabled modules (display=1): {len(enabled_rows)}')
